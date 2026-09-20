"""
Generates docs/TidalSql-User-Guide.docx.

This script is the source of truth for the Word document. Re-run it after editing to
regenerate the .docx:

    python docs/_generate_user_guide.py

Requires python-docx (pip install python-docx).
"""

import os
from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT

HERE = os.path.dirname(os.path.abspath(__file__))
OUT = os.path.join(HERE, "TidalSql-User-Guide.docx")

# ---- palette -------------------------------------------------------------------
INK = RGBColor(0x1B, 0x1F, 0x24)
ACCENT = RGBColor(0x0B, 0x5C, 0xAB)
MUTED = RGBColor(0x57, 0x60, 0x6A)
CODE_BG = "F3F5F7"
CODE_INK = RGBColor(0x0A, 0x2A, 0x12)


def _shade(cell, hex_fill):
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_fill)
    tc_pr.append(shd)


def code_block(doc, lines):
    """A single-cell shaded table used as a monospace code block."""
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    cell = table.cell(0, 0)
    _shade(cell, CODE_BG)
    cell.width = Inches(6.3)
    para_cell = cell.paragraphs[0]
    para_cell.paragraph_format.space_after = Pt(0)
    para_cell.paragraph_format.space_before = Pt(0)
    for i, line in enumerate(lines):
        run = para_cell.add_run(("\n" if i else "") + line)
        run.font.name = "Consolas"
        run.font.size = Pt(9.5)
        run.font.color.rgb = CODE_INK
    doc.add_paragraph().paragraph_format.space_after = Pt(4)


def code_inline(paragraph, text):
    run = paragraph.add_run(text)
    run.font.name = "Consolas"
    run.font.size = Pt(10)
    run.font.color.rgb = CODE_INK


def bullet(doc, runs):
    """runs: list of (text, is_code) tuples."""
    p = doc.add_paragraph(style="List Bullet")
    for text, is_code in runs:
        r = p.add_run(text)
        if is_code:
            r.font.name = "Consolas"
            r.font.size = Pt(10)
            r.font.color.rgb = CODE_INK
    return p


def para(doc, text, size=11, color=INK, bold=False, italic=False, after=6):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(after)
    r = p.add_run(text)
    r.font.size = Pt(size)
    r.font.color.rgb = color
    r.bold = bold
    r.italic = italic
    return p


def make_table(doc, headers, rows, widths=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Light Grid Accent 1"
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    hdr = table.rows[0].cells
    for i, h in enumerate(headers):
        hdr[i].paragraphs[0].add_run(h).bold = True
    code_prefixes = ("GRANT", "DENY", "REVOKE", "CREATE", "ALTER", "SELECT",
                     "INSERT", "UPDATE", "DELETE", "EXECUTE", "USE", "db_", "VIEW")
    for row in rows:
        cells = table.add_row().cells
        for i, val in enumerate(row):
            para_cell = cells[i].paragraphs[0]
            if val.startswith(code_prefixes):
                code_inline(para_cell, val)
            else:
                para_cell.add_run(val)
    if widths:
        for row in table.rows:
            for i, w in enumerate(widths):
                row.cells[i].width = Inches(w)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return table


# ================================================================================
doc = Document()

normal = doc.styles["Normal"]
normal.font.name = "Calibri"
normal.font.size = Pt(11)
normal.font.color.rgb = INK

# ---- Title --------------------------------------------------------------------
title = doc.add_paragraph()
tr = title.add_run("TidalSql")
tr.font.size = Pt(34)
tr.bold = True
tr.font.color.rgb = ACCENT

sub = doc.add_paragraph()
sr = sub.add_run("User Guide \u2014 Querying, Administration, and Security")
sr.font.size = Pt(15)
sr.font.color.rgb = MUTED

meta = doc.add_paragraph()
mr = meta.add_run("A T-SQL query engine over Parquet files")
mr.italic = True
mr.font.color.rgb = MUTED

para(doc,
     "TidalSql stores each table as one or more Apache Parquet files and lets you query and "
     "manage them with a familiar Transact-SQL dialect. It provides a SQL Server\u2013style "
     "security model (logins, users, roles, schemas, and GRANT/DENY/REVOKE), a web-based SQL "
     "editor with an object explorer, and an interactive console (REPL).",
     after=10)

doc.add_paragraph()

# ---- Contents -----------------------------------------------------------------
para(doc, "Contents", size=14, bold=True, color=ACCENT, after=4)
for item in [
    "1. Getting Started",
    "2. Core Concepts",
    "3. Working with Databases and Tables",
    "4. Querying Data",
    "5. Security: Logins, Users, Roles, and Permissions",
    "6. Security Recipes (Step-by-Step)",
    "7. The Table Overview (Web UI)",
    "8. Tips, Limits, and Troubleshooting",
    "9. Quick Reference",
]:
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(2)
    p.add_run(item).font.size = Pt(11)

doc.add_page_break()

# ================================================================================
doc.add_heading("1. Getting Started", level=1)
para(doc, "TidalSql ships as a .NET 8 solution with three ways to run it:", after=4)
make_table(doc,
    ["Component", "Project", "What it is"],
    [
        ["Engine library", "TidalSqlLib", "The T-SQL engine and security model. Referenced by the other apps."],
        ["Web app", "TidalSqlApi", "A browser SQL editor with a login screen and object explorer."],
        ["Console (REPL)", "TidalSqlConsole", "An interactive command-line session for running SQL and scripts."],
    ],
    widths=[1.7, 1.6, 3.0])

doc.add_heading("Run the console (REPL)", level=2)
code_block(doc, ["dotnet run --project TidalSqlConsole"])
para(doc, "On first run, TidalSql seeds a default administrator login and prints its "
          "credentials once. Sign in, then change the password.", after=4)
make_table(doc,
    ["Setting", "Value"],
    [
        ["Default admin login", "admin"],
        ["Initial password", "Printed once on first run (or set TIDALSQL_ADMIN_PASSWORD)"],
    ],
    widths=[2.4, 3.9])
para(doc, "Useful REPL commands:", after=2)
bullet(doc, [("Type any T-SQL statement and press Enter to execute it.", False)])
bullet(doc, [(":r <path>", True), ("  or  ", False), ("run <path>", True),
             ("  \u2014 execute a .sql script file (for example ", False),
             ("TidalScripts/build.sql", True), (").", False)])
bullet(doc, [("exit", True), ("  or  ", False), ("quit", True), ("  \u2014 end the session.", False)])

doc.add_heading("Run the web app", level=2)
code_block(doc, [
    "dotnet run --project TidalSqlApi",
    "# then open the printed URL, e.g. https://localhost:62549",
])
para(doc, "Sign in with your login on the web page. The left pane is an object explorer "
          "(databases, tables, procedures); the center is a SQL editor. Press Ctrl+Enter to run.",
     after=6)

# ================================================================================
doc.add_heading("2. Core Concepts", level=1)
make_table(doc,
    ["Concept", "How TidalSql implements it"],
    [
        ["Database", "A folder on disk. Create with CREATE DATABASE; select with USE."],
        ["Table", "A Parquet file (tables/<name>.parquet) or a directory of part files "
                  "(part-0.parquet \u2026 part-n.parquet) that form one logical table."],
        ["Schema", "A namespace for objects (dbo by default). Objects are addressed as [schema].[table]."],
        ["Metadata sidecar", "<table>.parquet.meta.json records identity, defaults, computed "
                             "columns, and nullability created by CREATE TABLE."],
        ["Login", "A server-level identity that authenticates with a password."],
        ["User", "A database-level principal mapped to a login. Access is granted to users/roles."],
    ],
    widths=[1.7, 4.6])
para(doc, "Multi-part (directory) tables are read-only \u2014 they are ideal for large, "
          "pre-generated data sets (for example TPC-H). Writes (INSERT/UPDATE/DELETE) target "
          "single-file tables.", italic=True, color=MUTED, after=6)

# ================================================================================
doc.add_heading("3. Working with Databases and Tables", level=1)
code_block(doc, [
    "CREATE DATABASE Sales;",
    "USE Sales;",
    "",
    "CREATE TABLE dbo.Customers (",
    "    Id        INT IDENTITY(1,1) NOT NULL,",
    "    Name      VARCHAR(100)      NOT NULL,",
    "    CreatedAt DATETIME2(6)      NOT NULL,",
    "    Credit    DECIMAL(12,2)     NULL);",
    "",
    "INSERT INTO Customers (Name, CreatedAt, Credit)",
    "VALUES ('Acme', SYSUTCDATETIME(), 500.00),",
    "       ('Globex', SYSUTCDATETIME(), NULL);",
])
para(doc, "List objects with the system catalog views:", after=2)
code_block(doc, [
    "SELECT name FROM sys.databases;",
    "SELECT name FROM sys.tables;",
    "SELECT name FROM sys.procedures;",
])

# ================================================================================
doc.add_heading("4. Querying Data", level=1)
para(doc, "TidalSql supports a broad T-SQL surface, including:", after=2)
for feat in [
    "SELECT with WHERE, GROUP BY, HAVING, ORDER BY, DISTINCT, and TOP (n).",
    "Joins (INNER/LEFT), subqueries, derived tables, and common table expressions (WITH).",
    "Set operators: UNION / UNION ALL, EXCEPT, INTERSECT.",
    "Aggregates (COUNT, SUM, MIN, MAX, AVG) and the ROW_NUMBER() window function.",
    "Scalar expressions: arithmetic, CASE, CAST/CONVERT, and ~25 built-in functions.",
    "Stored procedures (CREATE/ALTER PROCEDURE) and EXECUTE.",
]:
    bullet(doc, [(feat, False)])
para(doc, "Whole-table aggregates over a single table use a streaming fast path that reads "
          "only the referenced columns (projection pushdown) without materializing rows:",
     after=2, color=MUTED, italic=True)
code_block(doc, [
    "SELECT SUM(l_quantity) AS TotalQty, COUNT(*) AS Lines",
    "FROM   dbo.lineitem;",
])

# ================================================================================
doc.add_heading("5. Security: Logins, Users, Roles, and Permissions", level=1)
para(doc,
     "TidalSql mirrors SQL Server's authorization model. Authentication happens at the server "
     "with a login; authorization happens inside a database, where a login is mapped to a user "
     "that is granted permissions directly or through roles.",
     after=6)

doc.add_heading("The permission chain", level=2)
para(doc, "Login  \u2192  User (per database)  \u2192  Role membership  \u2192  GRANT/DENY on a securable.",
     bold=True, after=6)

doc.add_heading("Evaluation rules", level=2)
for rule in [
    "Role membership is expanded transitively; every principal is also a member of public.",
    "DENY overrides GRANT. A single DENY anywhere blocks the action.",
    "Permissions are inherited from database \u2192 schema \u2192 object (a broad grant cascades down).",
    "Ownership implies control: the owner of a schema controls its objects.",
    "sysadmin (server) and db_owner (database) bypass all permission checks.",
    "Default-deny: a mapped user with no grants has no access and does not even see the database.",
]:
    bullet(doc, [(rule, False)])

doc.add_heading("Server roles (fixed)", level=2)
make_table(doc,
    ["Server role", "Grants the ability to"],
    [
        ["sysadmin", "Do anything on the server and in every database (bypasses all checks)."],
        ["securityadmin", "Manage server security: create logins and manage server roles."],
        ["dbcreator", "Create new databases (CREATE DATABASE). The creator becomes db_owner of it."],
    ],
    widths=[1.7, 4.6])
para(doc, "Add a login to a server role with ALTER SERVER ROLE:", after=2)
code_block(doc, ["ALTER SERVER ROLE dbcreator ADD MEMBER alice_login;"])

doc.add_heading("Database roles (fixed)", level=2)
make_table(doc,
    ["Database role", "Implicit permission"],
    [
        ["db_owner", "Full control of the database (bypasses checks)."],
        ["db_securityadmin", "Manage users, roles, schemas, and permissions."],
        ["db_ddladmin", "Create and alter objects (DDL)."],
        ["db_datareader", "SELECT on all tables."],
        ["db_datawriter", "INSERT, UPDATE, DELETE on all tables."],
        ["db_executor", "EXECUTE on stored procedures."],
    ],
    widths=[1.9, 4.4])

doc.add_heading("Permissions and securables", level=2)
make_table(doc,
    ["Permission", "Meaning"],
    [
        ["SELECT", "Read rows from a table."],
        ["INSERT / UPDATE / DELETE", "Modify rows in a table."],
        ["EXECUTE", "Run a stored procedure."],
        ["ALTER", "Create/alter the object (DDL)."],
        ["VIEW DEFINITION", "See that the object exists (metadata visibility)."],
        ["CONTROL", "Full control of the securable; covers all of the above."],
    ],
    widths=[2.3, 4.0])
para(doc, "A permission can be granted at three securable scopes, from broad to specific:", after=2)
make_table(doc,
    ["Scope", "Syntax", "Applies to"],
    [
        ["Database", "GRANT SELECT TO <user>", "Every object in the current database."],
        ["Schema", "GRANT SELECT ON SCHEMA::sales TO <user>", "Every object in that schema."],
        ["Object", "GRANT SELECT ON dbo.Orders TO <user>", "One table (or procedure)."],
    ],
    widths=[1.2, 3.3, 1.8])

# ================================================================================
doc.add_heading("6. Security Recipes (Step-by-Step)", level=1)

doc.add_heading("Create a login and map a database user", level=2)
code_block(doc, [
    "-- Server scope: create the login (run as sysadmin/securityadmin).",
    "CREATE LOGIN alice WITH PASSWORD = 'S3cure!Pass';",
    "",
    "-- Database scope: map a user to the login.",
    "USE Sales;",
    "CREATE USER alice FOR LOGIN alice;",
])

doc.add_heading("Grant read-only access to one table", level=2)
code_block(doc, ["GRANT SELECT ON dbo.Orders TO alice;"])

doc.add_heading("Grant read access to the whole database", level=2)
code_block(doc, [
    "-- Option A: a fixed role.",
    "ALTER ROLE db_datareader ADD MEMBER alice;",
    "",
    "-- Option B: a database-scoped grant.",
    "GRANT SELECT TO alice;",
])

doc.add_heading("Create a custom role and assign members", level=2)
code_block(doc, [
    "CREATE ROLE SalesAnalyst;",
    "GRANT SELECT ON SCHEMA::dbo TO SalesAnalyst;",
    "GRANT EXECUTE ON dbo.RecalculateTotals TO SalesAnalyst;",
    "",
    "ALTER ROLE SalesAnalyst ADD MEMBER alice;",
    "ALTER ROLE SalesAnalyst ADD MEMBER bob;",
])

doc.add_heading("Grant write access", level=2)
code_block(doc, [
    "GRANT INSERT, UPDATE, DELETE ON dbo.Orders TO alice;",
    "-- or database-wide:",
    "ALTER ROLE db_datawriter ADD MEMBER alice;",
])

doc.add_heading("Explicitly deny (overrides any grant)", level=2)
code_block(doc, ["DENY SELECT ON dbo.SalaryData TO alice;"])

doc.add_heading("Revoke a previously granted or denied permission", level=2)
code_block(doc, ["REVOKE SELECT ON dbo.Orders FROM alice;"])

doc.add_heading("Let someone create databases", level=2)
code_block(doc, [
    "CREATE LOGIN carol WITH PASSWORD = 'An0ther!Pass';",
    "ALTER SERVER ROLE dbcreator ADD MEMBER carol;",
    "-- carol can now: CREATE DATABASE Marketing;  (she becomes db_owner of it)",
])

doc.add_heading("Inspect effective and object permissions", level=2)
bullet(doc, [("Web UI: open a table \u2192 Overview \u2192 ", False),
             ("Security and permissions", True),
             (" lists table/schema/database grants and denies.", False)])
bullet(doc, [("Web UI: the ", False), ("My Access", True),
             (" button shows the signed-in login's effective permissions.", False)])

# ================================================================================
doc.add_heading("7. The Table Overview (Web UI)", level=1)
para(doc, "Selecting a table in the explorer opens an Overview tab that summarizes the table "
          "and its access. The Summary card reports:", after=2)
for label, desc in [
    ("Rows / Columns / Size (on disk)", "current shape and footprint"),
    ("Storage", "single-file, or multi-part with the part count"),
    ("Created", "when the table's data was first written"),
    ("Last data change", "most recent write to the table data"),
    ("Last schema change", "when the table definition (metadata) last changed"),
    ("Last query run", "the last time a query read the table"),
]:
    bullet(doc, [(label + " \u2014 ", False), (desc, False)])
para(doc, "The Security and permissions section lists every GRANT/DENY that governs the table, "
          "tagged by scope (Table, Schema, or Database), so you can see exactly who can reach the "
          "data and why.", after=6)

# ================================================================================
doc.add_heading("8. Tips, Limits, and Troubleshooting", level=1)
make_table(doc,
    ["Situation", "What to do"],
    [
        ["\"Permission denied\" on SELECT",
         "The user has no covering grant. GRANT SELECT on the object/schema/database, or add "
         "them to db_datareader."],
        ["A user cannot see a database",
         "Access is default-deny. Map a user (CREATE USER ... FOR LOGIN) and grant at least one "
         "permission, or add them to a fixed role."],
        ["Cannot INSERT/UPDATE a table",
         "That table may be multi-part (directory of parts), which is read-only. Use a single-file "
         "table for writes."],
        ["\"requires sysadmin or dbcreator\"",
         "Only sysadmin or dbcreator members may CREATE DATABASE. Add the login to the dbcreator "
         "server role."],
        ["Forgot to select a database",
         "Run USE <database> first; most statements require a current database."],
        ["Managing security fails",
         "Creating logins needs sysadmin/securityadmin; creating users/roles and granting needs "
         "db_owner or db_securityadmin."],
    ],
    widths=[2.3, 4.0])

# ================================================================================
doc.add_heading("9. Quick Reference", level=1)
make_table(doc,
    ["Task", "Statement"],
    [
        ["Create database", "CREATE DATABASE <name>;"],
        ["Select database", "USE <name>;"],
        ["Create login", "CREATE LOGIN <name> WITH PASSWORD = '\u2026';"],
        ["Map user to login", "CREATE USER <name> FOR LOGIN <login>;"],
        ["Create role", "CREATE ROLE <name>;"],
        ["Add role member", "ALTER ROLE <role> ADD MEMBER <principal>;"],
        ["Remove role member", "ALTER ROLE <role> DROP MEMBER <principal>;"],
        ["Add server-role member", "ALTER SERVER ROLE <role> ADD MEMBER <login>;"],
        ["Create schema", "CREATE SCHEMA <name>;"],
        ["Grant on object", "GRANT SELECT ON <schema>.<table> TO <principal>;"],
        ["Grant on schema", "GRANT SELECT ON SCHEMA::<schema> TO <principal>;"],
        ["Grant database-wide", "GRANT SELECT TO <principal>;"],
        ["Deny", "DENY <permission> ON <object> TO <principal>;"],
        ["Revoke", "REVOKE <permission> ON <object> FROM <principal>;"],
        ["Impersonate / revert", "EXECUTE AS USER = '<user>';  \u2026  REVERT;"],
    ],
    widths=[2.3, 4.0])

para(doc, "", after=8)
foot = doc.add_paragraph()
fr = foot.add_run("Generated from docs/_generate_user_guide.py \u2014 re-run that script to update this document.")
fr.italic = True
fr.font.size = Pt(9)
fr.font.color.rgb = MUTED

doc.save(OUT)
print("wrote", OUT)
